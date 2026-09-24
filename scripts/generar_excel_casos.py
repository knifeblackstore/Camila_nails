# -*- coding: utf-8 -*-
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

wb = openpyxl.Workbook()

# ==========================================
# HOJA: MATRIZ DE CASOS DE PRUEBA (15 CASOS)
# ==========================================
ws1 = wb.active
ws1.title = "Casos de Prueba (15)"
ws1.views.sheetView[0].showGridLines = True

# Paleta de colores institucional SENA
verde_sena = "39A900"
verde_oscuro = "1E5B00"
verde_claro = "E8F5E9"
gris_claro = "F4F6F7"
borde_gris = "D0D3D4"

font_titulo = Font(name="Calibri", size=13, bold=True, color="FFFFFF")
font_subtitulo = Font(name="Calibri", size=11, bold=True, color="1E5B00")
font_header_tabla = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
font_bold = Font(name="Calibri", size=10, bold=True, color="000000")
font_normal = Font(name="Calibri", size=9.5, color="202020")
font_aprobado = Font(name="Calibri", size=10, bold=True, color="1B5E20")

fill_verde_sena = PatternFill(start_color=verde_sena, end_color=verde_sena, fill_type="solid")
fill_verde_oscuro = PatternFill(start_color=verde_oscuro, end_color=verde_oscuro, fill_type="solid")
fill_verde_claro = PatternFill(start_color=verde_claro, end_color=verde_claro, fill_type="solid")
fill_gris = PatternFill(start_color=gris_claro, end_color=gris_claro, fill_type="solid")
fill_aprobado = PatternFill(start_color="C8E6C9", end_color="C8E6C9", fill_type="solid")

thin_border = Border(
    left=Side(style='thin', color=borde_gris),
    right=Side(style='thin', color=borde_gris),
    top=Side(style='thin', color=borde_gris),
    bottom=Side(style='thin', color=borde_gris)
)

# Encabezado General
ws1.merge_cells("A1:K1")
ws1["A1"] = "SERVICIO NACIONAL DE APRENDIZAJE - SENA | REGIONAL ANTIOQUIA"
ws1["A1"].font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
ws1["A1"].fill = fill_verde_oscuro
ws1["A1"].alignment = Alignment(horizontal="center", vertical="center")
ws1.row_dimensions[1].height = 24

ws1.merge_cells("A2:K2")
ws1["A2"] = "EVIDENCIA GA9-220501096-AA2-EV01: DISEÑA CASOS Y DEFINE EL AMBIENTE DE PRUEBAS DE SOFTWARE SEGÚN PROYECTO"
ws1["A2"].font = font_titulo
ws1["A2"].fill = fill_verde_sena
ws1["A2"].alignment = Alignment(horizontal="center", vertical="center")
ws1.row_dimensions[2].height = 28

# Metadatos del Aprendiz y Proyecto
metadatos = [
    ("Programa:", "Análisis y Desarrollo de Software (ADSO)", "Ficha:", "3186645"),
    ("Aprendiz / Analista:", "Andrés Mauricio Valencia Arango", "Instructor:", "Adornay Sanchez"),
    ("Proyecto:", "Sistema de Gestión Web 'Camila Nails' (React + Node.js + MySQL)", "Total de Casos:", "15 Casos Diseñados y Aprobados")
]

row_idx = 4
for label1, val1, label2, val2 in metadatos:
    ws1[f"A{row_idx}"] = label1
    ws1[f"A{row_idx}"].font = font_bold
    ws1[f"A{row_idx}"].fill = fill_verde_claro
    
    ws1.merge_cells(f"B{row_idx}:F{row_idx}")
    ws1[f"B{row_idx}"] = val1
    ws1[f"B{row_idx}"].font = font_normal
    
    ws1[f"G{row_idx}"] = label2
    ws1[f"G{row_idx}"].font = font_bold
    ws1[f"G{row_idx}"].fill = fill_verde_claro
    
    ws1.merge_cells(f"H{row_idx}:K{row_idx}")
    ws1[f"H{row_idx}"] = val2
    ws1[f"H{row_idx}"].font = font_normal
    
    for col in range(1, 12):
        col_letter = get_column_letter(col)
        ws1[f"{col_letter}{row_idx}"].border = thin_border
    
    ws1.row_dimensions[row_idx].height = 20
    row_idx += 1

# Sección: Definición del Ambiente de Pruebas
row_idx += 1
ws1.merge_cells(f"A{row_idx}:K{row_idx}")
ws1[f"A{row_idx}"] = "1. DEFINICIÓN DEL AMBIENTE DE PRUEBAS DE SOFTWARE"
ws1[f"A{row_idx}"].font = font_titulo
ws1[f"A{row_idx}"].fill = fill_verde_oscuro
ws1[f"A{row_idx}"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
ws1.row_dimensions[row_idx].height = 24
row_idx += 1

ambiente_datos = [
    ("Componente", "Detalle Técnico y Configuración del Ambiente"),
    ("Sistema Operativo", "Windows 11 Home / Pro (Arquitectura 64-bit)"),
    ("Servidor Backend", "Node.js v18+ con Express.js escuchando en http://localhost:4000"),
    ("Base de Datos", "MySQL Server 8.0 en puerto 3306 (Base: camila_nails con motor InnoDB)"),
    ("Cliente Frontend", "React.js empaquetado con Vite ejecutándose en http://localhost:5173"),
    ("Navegadores", "Google Chrome v128+ y Microsoft Edge (Pruebas de visualización y responsividad)"),
    ("Herramientas de Prueba", "Postman v10+ (Pruebas de API REST) y Chrome DevTools (Inspección UI / CSS)")
]

for item, desc in ambiente_datos:
    is_hdr = (item == "Componente")
    ws1.merge_cells(f"A{row_idx}:C{row_idx}")
    ws1[f"A{row_idx}"] = item
    ws1[f"A{row_idx}"].font = font_header_tabla if is_hdr else font_bold
    ws1[f"A{row_idx}"].fill = fill_verde_sena if is_hdr else fill_verde_claro
    ws1[f"A{row_idx}"].alignment = Alignment(horizontal="center" if is_hdr else "left", vertical="center", indent=1)
    
    ws1.merge_cells(f"D{row_idx}:K{row_idx}")
    ws1[f"D{row_idx}"] = desc
    ws1[f"D{row_idx}"].font = font_header_tabla if is_hdr else font_normal
    ws1[f"D{row_idx}"].fill = fill_verde_sena if is_hdr else PatternFill(fill_type=None)
    ws1[f"D{row_idx}"].alignment = Alignment(horizontal="center" if is_hdr else "left", vertical="center", indent=1)
    
    for col in range(1, 12):
        col_letter = get_column_letter(col)
        ws1[f"{col_letter}{row_idx}"].border = thin_border
        
    ws1.row_dimensions[row_idx].height = 20
    row_idx += 1

# Sección: Formato de Casos de Prueba (15 Casos)
row_idx += 1
ws1.merge_cells(f"A{row_idx}:K{row_idx}")
ws1[f"A{row_idx}"] = "2. MATRIZ DE CASOS DE PRUEBA (15 CASOS DETALLADOS SEGÚN PROYECTO)"
ws1[f"A{row_idx}"].font = font_titulo
ws1[f"A{row_idx}"].fill = fill_verde_oscuro
ws1[f"A{row_idx}"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
ws1.row_dimensions[row_idx].height = 24
row_idx += 1

headers_columnas = [
    ("N° Caso", 10),
    ("Herramienta Utilizada", 16),
    ("Autor del Caso de Prueba", 24),
    ("Descripción / Módulo", 28),
    ("Salida Esperada", 30),
    ("Salida Obtenida", 30),
    ("Resultado", 14),
    ("Seguimiento", 14),
    ("Severidad", 12),
    ("Evidencia", 28),
    ("Firma de Aprobación", 24)
]

for col_i, (h_title, _) in enumerate(headers_columnas, start=1):
    cell = ws1.cell(row=row_idx, column=col_i, value=h_title)
    cell.font = font_header_tabla
    cell.fill = fill_verde_sena
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = thin_border

ws1.row_dimensions[row_idx].height = 32
row_idx += 1

autor_firma = "Andrés Mauricio Valencia Arango"

casos_15 = [
    (
        "CP-001",
        "Postman",
        autor_firma,
        "Registro exitoso de nuevo cliente (POST /api/register)",
        "Código HTTP 200 OK con JSON {'success': true, 'message': 'Registro exitoso'}",
        "Código HTTP 200 OK recibido y registro almacenado en tabla 'users' de MySQL",
        "Aprobado",
        "N/A",
        "Baja",
        "Respuesta JSON status 200 en Postman y persistencia en DB",
        autor_firma
    ),
    (
        "CP-002",
        "Postman",
        autor_firma,
        "Validación de campos obligatorios en registro (POST /api/register sin email)",
        "Código HTTP 400 Bad Request con mensaje {'error': 'Faltan campos'}",
        "Código HTTP 400 Bad Request devuelto, bloqueando inserción incompleta",
        "Aprobado",
        "N/A",
        "Media",
        "Captura de pantalla en Postman con respuesta 400 Bad Request",
        autor_firma
    ),
    (
        "CP-003",
        "Postman",
        autor_firma,
        "Control de usuario duplicado en registro (POST /api/register con email existente)",
        "Código HTTP 409 Conflict con mensaje {'error': 'Usuario o email ya existe'}",
        "Código HTTP 409 Conflict recibido impidiendo duplicidad de cuentas",
        "Aprobado",
        "N/A",
        "Media",
        "Validación de restricción UNIQUE en MySQL retornando 409 en Postman",
        autor_firma
    ),
    (
        "CP-004",
        "Postman",
        autor_firma,
        "Autenticación exitosa de usuario administrador (POST /api/login)",
        "Código HTTP 200 OK con datos de perfil y validación de rol 'admin'",
        "Acceso otorgado con status 200 OK y confirmación de rol administrativo",
        "Aprobado",
        "N/A",
        "Alta",
        "Contraseña validada con hash bcryptjs y retorno de sesión en Postman",
        autor_firma
    ),
    (
        "CP-005",
        "Postman",
        autor_firma,
        "Rechazo de inicio de sesión con contraseña errada (POST /api/login)",
        "Código HTTP 401 Unauthorized con mensaje {'error': 'Contraseña incorrecta'}",
        "Código HTTP 401 recibido, impidiendo accesos no autorizados al sistema",
        "Aprobado",
        "N/A",
        "Alta",
        "Respuesta HTTP 401 Unauthorized documentada en Postman",
        autor_firma
    ),
    (
        "CP-006",
        "Postman",
        autor_firma,
        "Consulta general de usuarios registrados (GET /api/users)",
        "Arreglo JSON con usuarios y campos públicos excluyendo contraseñas",
        "Respuesta 200 OK con lista completa de usuarios protegida sin hash",
        "Aprobado",
        "N/A",
        "Baja",
        "Array de usuarios en formato JSON retornado con Status 200 OK",
        autor_firma
    ),
    (
        "CP-007",
        "Postman",
        autor_firma,
        "Modificación de rol de usuario (PATCH /api/users/:id/role)",
        "Cambio de rol de 'user' a 'admin' reflejado en base de datos con status 200 OK",
        "Rol actualizado exitosamente y verificado con consulta posterior",
        "Aprobado",
        "N/A",
        "Media",
        "Respuesta PATCH status 200 OK con confirmación de cambio de rol",
        autor_firma
    ),
    (
        "CP-008",
        "Postman / Chrome",
        autor_firma,
        "Edición de datos de usuario en panel admin (PUT /api/users/:id)",
        "Actualización de nombre y correo electrónico con confirmación 200 OK",
        "Datos actualizados en MySQL y reflejados en el formulario en vivo",
        "Aprobado",
        "N/A",
        "Media",
        "Respuesta PUT 200 OK {'message': 'Usuario actualizado'}",
        autor_firma
    ),
    (
        "CP-009",
        "Postman / Chrome",
        autor_firma,
        "Eliminación física de usuario desde panel (DELETE /api/users/:id)",
        "Eliminación del registro con confirmación previa y respuesta 200 OK",
        "Status 200 OK {'message': 'Usuario eliminado'} y registro removido",
        "Aprobado",
        "N/A",
        "Alta",
        "Comprobación de registro ausente tras petición DELETE ejecutada",
        autor_firma
    ),
    (
        "CP-010",
        "Postman / Chrome",
        autor_firma,
        "Creación manual de Ficha de Cliente desde Panel (POST /api/fichas)",
        "Inserción en tabla 'fichas' con ID autoincremental y status 200 OK",
        "Ficha creada con éxito, status 200 OK y asignación de insertId",
        "Aprobado",
        "N/A",
        "Alta",
        "Ficha persistida en MySQL y visible de inmediato en la lista del panel",
        autor_firma
    ),
    (
        "CP-011",
        "Postman",
        autor_firma,
        "Consulta y listado histórico de Fichas Clínicas (GET /api/fichas)",
        "Retorno de todas las fichas ordenadas cronológicamente con status 200",
        "Arreglo JSON recibido con todos los historiales y servicios registrados",
        "Aprobado",
        "N/A",
        "Baja",
        "Respuesta GET 200 OK visualizada en Postman con arreglo de fichas",
        autor_firma
    ),
    (
        "CP-012",
        "Postman / Chrome",
        autor_firma,
        "Actualización y edición de Ficha de Cliente (PUT /api/fichas/:id)",
        "Modificación de teléfono y servicio asignado con status 200 OK",
        "Campos modificados en MySQL con mensaje {'message': 'Ficha actualizada'}",
        "Aprobado",
        "N/A",
        "Media",
        "Petición PUT 200 OK en Postman y sincronización en interfaz React",
        autor_firma
    ),
    (
        "CP-013",
        "Postman / Chrome",
        autor_firma,
        "Eliminación física de Ficha de Cliente (DELETE /api/fichas/:id)",
        "Registro de ficha eliminado de la base de datos con status 200 OK",
        "Status 200 OK {'message': 'Ficha eliminada'} y registro borrado en DB",
        "Aprobado",
        "N/A",
        "Media",
        "Comprobación en tabla 'fichas' y actualización de la lista en React",
        autor_firma
    ),
    (
        "CP-014",
        "Chrome DevTools",
        autor_firma,
        "Contraste y legibilidad de inputs en Modo Oscuro (Formularios Admin)",
        "El texto debe ser claramente legible (blanco sobre fondo oscuro) al escribir",
        "Se detectó fondo blanco en focus que ocultaba la letra; fue corregido en App.css",
        "Aprobado",
        "BUG-UI-01",
        "Media",
        "Regla CSS modificada (App.css) e inspección exitosa en DevTools",
        autor_firma
    ),
    (
        "CP-015",
        "Chrome DevTools",
        autor_firma,
        "Diseño responsivo del Home y Navbar en móviles (360px a 412px)",
        "Adaptación fluida de lado a lado sin marco rígido ni desbordamiento lateral",
        "Ancho fijo de 1126px reemplazado por 100% fluido con padding móvil perfecto",
        "Aprobado",
        "N/A",
        "Baja",
        "Emulación responsive en Google Chrome DevTools con visual amplia",
        autor_firma
    )
]

for cp in casos_15:
    c_num, c_tool, c_autor, c_desc, c_esp, c_obt, c_res, c_seg, c_sev, c_evi, c_firma = cp
    
    ws1.cell(row=row_idx, column=1, value=c_num).alignment = Alignment(horizontal="center", vertical="center")
    ws1.cell(row=row_idx, column=2, value=c_tool).alignment = Alignment(horizontal="center", vertical="center")
    ws1.cell(row=row_idx, column=3, value=c_autor).alignment = Alignment(horizontal="left", vertical="center")
    ws1.cell(row=row_idx, column=4, value=c_desc).alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    ws1.cell(row=row_idx, column=5, value=c_esp).alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    ws1.cell(row=row_idx, column=6, value=c_obt).alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    
    res_cell = ws1.cell(row=row_idx, column=7, value=c_res)
    res_cell.alignment = Alignment(horizontal="center", vertical="center")
    res_cell.font = font_aprobado
    res_cell.fill = fill_aprobado
    
    ws1.cell(row=row_idx, column=8, value=c_seg).alignment = Alignment(horizontal="center", vertical="center")
    ws1.cell(row=row_idx, column=9, value=c_sev).alignment = Alignment(horizontal="center", vertical="center")
    ws1.cell(row=row_idx, column=10, value=c_evi).alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    
    firma_cell = ws1.cell(row=row_idx, column=11, value=c_firma)
    firma_cell.alignment = Alignment(horizontal="center", vertical="center")
    firma_cell.font = font_bold
    
    for col in range(1, 12):
        col_letter = get_column_letter(col)
        c = ws1[f"{col_letter}{row_idx}"]
        c.border = thin_border
        if col != 7:
            if col != 11:
                c.font = font_normal
            if row_idx % 2 == 0 and col != 11:
                c.fill = fill_gris
                
    ws1.row_dimensions[row_idx].height = 42
    row_idx += 1

# Ajustar ancho de columnas
for col_i, (_, w) in enumerate(headers_columnas, start=1):
    col_letter = get_column_letter(col_i)
    ws1.column_dimensions[col_letter].width = w

# Guardar libro Excel
excel_path = "Evidencia_GA9-220501096-AA2-EV01_CasosPrueba.xlsx"
wb.save(excel_path)
print(f"Archivo Excel con 15 casos generado exitosamente: {excel_path}")
